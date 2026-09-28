from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ACCENT = RGBColor(0x1A, 0x1A, 0x1A)

def base(title):
    d = Document()
    st = d.styles['Normal']
    st.font.name = 'Georgia'; st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(8)
    h = d.add_heading(title, level=0)
    return d

def h2(d, text):
    h = d.add_heading(text, level=1)
    for r in h.runs: r.font.size = Pt(13)
    return h

def p(d, text, bold_prefix=None):
    par = d.add_paragraph()
    if bold_prefix:
        r = par.add_run(bold_prefix); r.bold = True
    par.add_run(text)
    return par

# ---------- Terms ----------
d = base('Rankly – Terms of Use')
p(d, 'Last updated: 21 September 2026. Rankly is a demonstration pay-to-rank product leaderboard. These terms explain how the demo works and what would apply in a production version.')
h2(d, '1. What Rankly is')
p(d, 'Rankly lets products bid for placement on a public leaderboard. Higher bids earn higher ranks. The current build is a demo: bids are illustrative, checkout is a test preview, and no real payment is processed.')
h2(d, '2. Bids and placement')
p(d, 'A bid is an offer for rank position, not a purchase of goods or services. Rank is determined by bid amount relative to other live bids. Rankly does not guarantee any placement, traffic, or outcome from a bid.')
h2(d, '3. Payments (production)')
p(d, 'In production, payments would be processed by Razorpay (UPI, cards, netbanking). A bid is captured only after successful payment. Failed or cancelled payments create no rank claim.')
h2(d, '4. Refunds')
p(d, 'Demo bids are not charged and cannot be refunded. A production version would publish a refund policy covering duplicate charges, failed placements, and bid withdrawals before publishing this page.')
h2(d, '5. Acceptable listings')
p(d, 'Listings must be real products with accurate descriptions. Prohibited: illegal goods or services, deceptive claims, malware, hate content, and anything infringing intellectual property. Rankly may remove listings that violate these rules without refund of the bid.')
h2(d, '6. Fair use')
p(d, 'Do not manipulate rankings through fake bids, automated submissions, or payment fraud. Violations may lead to removal and a ban from future bidding.')
h2(d, '7. Changes')
p(d, 'Rankly may update these terms as the product evolves. Continued use after changes means you accept the updated terms.')
d.save('/home/hatch/workspace/rankly/docs/rankly-terms.docx')

# ---------- Privacy ----------
d = base('Rankly – Privacy Policy')
p(d, 'Last updated: 21 September 2026. This policy describes what data Rankly collects in the demo and what a production version would collect.')
h2(d, '1. Demo data handling')
p(d, 'The demo runs entirely in your browser. Form entries (product name, category, bid) are not sent to any server and are not persisted. Closing or reloading the page discards them.')
h2(d, '2. Production data collection')
p(d, 'A production version would collect: ', bold_prefix=None)
for item in ['Contact details you provide when claiming a rank (name, email).',
             'Payment details processed securely by Razorpay — Rankly never sees or stores card numbers.',
             'Listing details: product name, URL, category, description, and bid history.',
             'Basic analytics: page views and clicks used to compute leaderboard statistics.']:
    d.add_paragraph(item, style='List Bullet')
h2(d, '3. How data is used')
p(d, 'Data is used to operate the leaderboard: display rankings, process bids, prevent fraud, and publish aggregate statistics. We do not sell personal data.')
h2(d, '4. Data retention')
p(d, 'Production records would be kept while your listing is active and for a reasonable period afterwards for accounting and fraud prevention, then deleted or anonymized.')
h2(d, '5. Your rights')
p(d, 'You may request a copy, correction, or deletion of your personal data at any time by contacting the site operator.')
h2(d, '6. Cookies')
p(d, 'The demo sets no cookies. A production version would disclose any cookies used for sessions, payments, or analytics before collecting payment.')
d.save('/home/hatch/workspace/rankly/docs/rankly-privacy.docx')
print('built')
